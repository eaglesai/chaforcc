"""IO helper utilities."""
from __future__ import annotations

import csv
import json
import logging
from pathlib import Path
from typing import Iterator, List, Mapping

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - optional dependency
    load_workbook = None  # type: ignore

LOGGER = logging.getLogger(__name__)


def read_seed_file(path: Path) -> List[Mapping[str, str]]:
    """Read a CSV or XLSX seed file and return rows as dictionaries."""

    if path.suffix.lower() == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as f:
            reader = csv.DictReader(f)
            return [row for row in reader]
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        if load_workbook is None:
            raise RuntimeError("openpyxl is required to parse Excel seed files")
        workbook = load_workbook(path, read_only=True)
        sheet = workbook.active
        headers: List[str] = []
        rows: List[Mapping[str, str]] = []
        for row_index, row in enumerate(sheet.iter_rows(values_only=True)):
            if row_index == 0:
                headers = [str(value).strip() if value is not None else "" for value in row]
                continue
            record = {headers[i]: ("" if value is None else str(value).strip()) for i, value in enumerate(row)}
            rows.append(record)
        return rows
    raise ValueError(f"Unsupported seed file extension: {path.suffix}")


def write_json(path: Path, data: Mapping) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)


def read_json(path: Path) -> Mapping:
    with path.open("r", encoding="utf-8") as f:
        return json.load(f)


def iter_artifact_directories(base_dir: Path) -> Iterator[Path]:
    for child in sorted(base_dir.iterdir()):
        if child.is_dir():
            yield child


def iter_json_files(base_dir: Path) -> Iterator[Path]:
    for path in sorted(base_dir.rglob("*.json")):
        if path.name == "manifest.json" or path.name.endswith("_report.json"):
            continue
        yield path
